# -*- coding: utf-8 -*-
"""产物验收：只输出 ASCII 布尔/计数，避免控制台编码干扰。"""
import json, sys
from docx import Document
from docx.oxml.ns import qn

path = sys.argv[1]
doc = Document(path)
paras = [(p.style.name, p.text) for p in doc.paragraphs]
full = "\n".join(p[1] for p in paras)
styles = sorted({s0 for s0, _ in paras})

checks = {
    "open_ok": True,
    "paragraph_count": len(paras),
    "styles_all_present": all(s in styles for s in ["Title", "Heading 1", "Heading 2", "Heading 3", "List Paragraph", "Quote"]),
    "table_count_2": len(doc.tables) == 2,
    "overview_table_rows_4": len(doc.tables[0].rows) == 4,
    "markdown_table_parsed": [c.text for c in doc.tables[1].rows[1].cells] == ["外壳长度", "3.65", "3.14"],
    "bullet_present": any(p[1].startswith("• 外壳长度：") for p in paras),
    "ordered_present": any(p[1].startswith("1. 优先关注") for p in paras),
    "no_markdown_leftover": not any(("###" in p[1]) or ("**" in p[1]) or ("|---" in p[1]) for p in paras),
    "xml_escape_ok": "不良统计 & 特殊字符 <测试>" in full,
    "not_done_listed": "未生成 AI 分析：本次未导入不良记录，无法生成柏拉图。" in full,
    "appendix_present": any("附：未生成分析的模块与原因" in p[1] for p in paras),
    "disclaimer_present": "不替代质量判定与工程评审" in full,
    "heading1_first_is_overview": [p[1] for p in paras if p[0] == "Heading 1"][0] == "模块总览",
    "module_numbering": [p[1] for p in paras if p[0] == "Heading 1"][1].startswith("一、"),
}
checks["ALL_PASS"] = all(v is True for k, v in checks.items() if isinstance(v, bool))
print(json.dumps(checks, ensure_ascii=True, indent=1))