# -*- coding: utf-8 -*-
"""验收真浏览器下载到的 .docx（桩 LLM 固定正文）。"""
import json, sys
from docx import Document

path = sys.argv[1]
doc = Document(path)
paras = [(p.style.name, p.text) for p in doc.paragraphs]
full = "\n".join(t for _, t in paras)
styles = {s for s, _ in paras}
tables = [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]

checks = {
    "open_ok": True,
    "has_title": any(s == "Title" for s, _ in paras),
    "has_heading1": "Heading 1" in styles,
    "has_heading2": "Heading 2" in styles,
    "has_heading3": "Heading 3" in styles,
    "has_list_paragraph": "List Paragraph" in styles,
    "table_count": len(tables),
    "overview_table_3cols": len(tables[0][0]) == 3 if tables else False,
    "markdown_table_row_parsed": any(r == ["外壳长度", "3.65"] for t in tables for r in t),
    "stub_text_present": "桩服务" in full,
    "bullet_present": any(t.startswith("• 复核实测数据") for _, t in paras),
    "ordered_present": any(t.startswith("1. 收紧上公差") for _, t in paras),
    "no_markdown_leftover": not any(("###" in t or "**" in t or "|---" in t) for _, t in paras),
    "disclaimer_present": "不替代质量判定与工程评审" in full,
}
ok_flags = [v for v in checks.values() if isinstance(v, bool)]
checks["ALL_PASS"] = all(ok_flags)
print(json.dumps(checks, ensure_ascii=True, indent=1))