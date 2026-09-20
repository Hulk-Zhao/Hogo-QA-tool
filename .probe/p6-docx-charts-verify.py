# -*- coding: utf-8 -*-
"""验收真浏览器下载到的 .docx 是否真的内嵌了图表 PNG（P6-A）。"""
import json, re, sys, zipfile
from docx import Document

path = sys.argv[1]
CONTENT_W = 5940000
MAX_H = 7200000

doc = Document(path)
paras = [(p.style.name, p.text) for p in doc.paragraphs]
full = "\n".join(t for _, t in paras)
styles = {s for s, _ in paras}
tables = [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
shapes = doc.inline_shapes

zf = zipfile.ZipFile(path)
names = zf.namelist()
media = sorted(n for n in names if n.startswith("word/media/"))
document_xml = zf.read("word/document.xml").decode("utf-8")
rels = zf.read("word/_rels/document.xml.rels").decode("utf-8")
types_xml = zf.read("[Content_Types].xml").decode("utf-8")

png_info = []
for n in media:
    data = zf.read(n)
    png_info.append({"name": n, "bytes": len(data), "is_png": data[:8] == b"\x89PNG\r\n\x1a\n"})

shape_dims = []
for s in shapes:
    w = int(s.width)
    h = int(s.height)
    shape_dims.append({"w": w, "h": h, "in_box": 0 < w <= CONTENT_W and 0 < h <= MAX_H})

rel_ids = re.findall(r'Id="([^"]+)"', rels)

checks = {
    "open_ok": True,
    "has_title": any(s == "Title" for s, _ in paras),
    "has_heading1": "Heading 1" in styles,
    "overview_table_3cols": len(tables[0][0]) == 3 if tables else False,
    "stub_text_present": "桩服务" in full,
    "chart_heading_present": any(t.strip() == "图表" and s == "Heading 1" for s, t in paras),
    "inline_shape_count": len(shapes),
    "media_count": len(media),
    "media_all_png": len(media) > 0 and all(p["is_png"] for p in png_info),
    "media_not_tiny": len(media) > 0 and all(p["bytes"] > 2000 for p in png_info),
    "media_matches_shapes": len(media) == len(shapes),
    "shapes_inside_page": len(shapes) > 0 and all(d["in_box"] for d in shape_dims),
    "document_xml_has_xmlns_r": "xmlns:r=" in document_xml,
    "document_xml_has_embed": 'r:embed="rIdImg' in document_xml,
    "rels_has_image": "/relationships/image" in rels,
    "rel_ids_unique": len(set(rel_ids)) == len(rel_ids),
    "content_types_has_png": 'Extension="png"' in types_xml,
    "no_markdown_leftover": not any(("###" in t or "**" in t or "|---" in t) for _, t in paras),
    "disclaimer_present": "不替代质量判定与工程评审" in full,
    "shape_dims": shape_dims,
    "png_info": png_info,
}
flags = [v for k, v in checks.items() if isinstance(v, bool) and k != "ALL_PASS"]
checks["ALL_PASS"] = all(flags)
print(json.dumps(checks, ensure_ascii=True, indent=1))